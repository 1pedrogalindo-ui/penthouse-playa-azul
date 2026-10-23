from openpyxl import load_workbook
from datetime import datetime
from pathlib import Path
import sys

source=Path(sys.argv[1]); target=Path(sys.argv[2]); wb=load_workbook(source,data_only=True,read_only=True)
def q(v):
    if v is None: return 'null'
    if isinstance(v,(int,float)): return str(v)
    if isinstance(v,datetime): v=v.date().isoformat()
    return "'"+str(v).replace("'","''")+"'"
lines=["begin;","truncate public.accounting_transactions, public.bookings, public.expenses, public.property_dues, public.barters restart identity cascade;"]
for row in list(wb['Movimientos'].iter_rows(min_row=4,max_row=51,values_only=True)):
    d,t,desc,typ,doc,benef,amount,balance,cat,note=row
    lines.append("insert into public.accounting_transactions(date,time,description,type,document_no,beneficiary,amount,bank_balance,category,notes,source) values("+','.join(map(q,[d,t,desc,typ,doc,benef,amount,balance,cat,note,'Excel Contabilidad_PH1401']))+");")
for row in list(wb['Ingresos hospedaje'].iter_rows(min_row=5,max_row=11,values_only=True)):
    paid,name,amount,stay,note=row
    lines.append("insert into public.bookings(date,guest_name,check_in,check_out,guests,season,nightly_rate,total,deposit,status,payment_status,source,notes) values("+','.join(map(q,[paid,name,stay,stay.replace(day=min(stay.day+1,28)) if isinstance(stay,datetime) else stay,1,'mid',amount,amount,amount,'completed','paid','Excel',note]))+");")
for row in list(wb['Facturas y gastos'].iter_rows(min_row=5,max_row=9,values_only=True)):
    d,supplier,taxid,invoice,buyer,desc,base,tax,total,statement,phbase,phtax,phtotal,note=row
    lines.append("insert into public.expenses(date,supplier,supplier_tax_id,invoice_no,description,base_amount,tax_amount,total,category,in_bank_statement,notes) values("+','.join(map(q,[d,supplier,taxid,invoice,desc,phbase,phtax,phtotal,'Factura PH',statement not in ('No encontrado','No encontrado (efectivo)','No (tarjeta de crédito)','No encontrado (forma: sin sistema financiero)'),note]))+");")
for row in list(wb['Deymar PH 1401'].iter_rows(min_row=6,max_row=11,values_only=True)):
    period,start,end,fee,water,total,note=row
    if fee: lines.append("insert into public.property_dues(date,period,kind,amount,status,notes) values("+','.join(map(q,[start,period,'Alícuota',fee,'paid',note]))+");")
    if water: lines.append("insert into public.property_dues(date,period,kind,amount,status,notes) values("+','.join(map(q,[start,period,'Agua',water,'paid',note]))+");")
for row in list(wb['Canjes'].iter_rows(min_row=5,max_row=5,values_only=True)):
    d,desc,party,amount,note=row; lines.append("insert into public.barters(date,description,counterparty,amount,notes) values("+','.join(map(q,row))+ ");")
lines.append('commit;'); target.write_text('\n'.join(lines),encoding='utf-8')
